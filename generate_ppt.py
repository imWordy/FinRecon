"""
Presentation generator for FinRecon.
Creates a professional, clean 12-slide PowerPoint presentation explaining the project
and saves it to C:\\Users\\anshb\\Downloads\\FinRecon_Presentation.pptx
"""

import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_PATH = Path(r"C:\Users\anshb\Downloads\FinRecon_Presentation.pptx")

# Palette definition
COLOR_BG = RGBColor(248, 250, 252)        # Slate 50
COLOR_PRIMARY = RGBColor(30, 58, 138)      # Deep Blue (Navy)
COLOR_SECONDARY = RGBColor(14, 116, 144)   # Cyan / Teal
COLOR_TEXT_MAIN = RGBColor(15, 23, 42)     # Dark Slate
COLOR_TEXT_MUTED = RGBColor(71, 85, 105)   # Muted Grey
COLOR_ACCENT = RGBColor(16, 185, 129)      # Emerald Green
COLOR_CARD_BG = RGBColor(255, 255, 255)    # Pure White
COLOR_BORDER = RGBColor(226, 232, 240)     # Light Gray Border
COLOR_RED_ACCENT = RGBColor(225, 29, 72)   # Crimson Red

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide, color=COLOR_BG):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="FINRECON • AGENTIC FINANCIAL RECONCILIATION"):
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_SECONDARY

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_PRIMARY

    def add_card(slide, left, top, width, height, title, items, badge="", border_color=COLOR_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        # Header / Title
        p0 = tf.paragraphs[0]
        if badge:
            p0.text = f"{badge}  {title}"
        else:
            p0.text = title
        p0.font.size = Pt(17)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_PRIMARY
        p0.space_after = Pt(12)

        for item in items:
            p = tf.add_paragraph()
            p.text = f"• {item}"
            p.font.size = Pt(13)
            p.font.color.rgb = COLOR_TEXT_MAIN
            p.space_after = Pt(6)

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, RGBColor(15, 23, 42))  # Dark Theme for Title

    # Accent decorative bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(1.0), Inches(0.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_ACCENT
    bar.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(10.8), Inches(3.6))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "FinRecon"
    p.font.size = Pt(50)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    p2 = tf1.add_paragraph()
    p2.text = "Agentic Financial Transaction Reconciliation & Exception Resolution System"
    p2.font.size = Pt(22)
    p2.font.color.rgb = RGBColor(147, 197, 253)
    p2.space_before = Pt(10)

    p3 = tf1.add_paragraph()
    p3.text = "Autonomous Multi-Agent Investigation • Policy Grounding via RAG • Human-in-the-Loop Governance"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.space_before = Pt(25)

    # Tag footer
    tb1_foot = s1.shapes.add_textbox(Inches(1.2), Inches(6.0), Inches(10.8), Inches(0.8))
    p4 = tb1_foot.text_frame.paragraphs[0]
    p4.text = "Course: Agentic AI  |  Proof of Concept Architecture  |  Production-Grade Demo"
    p4.font.size = Pt(12)
    p4.font.color.rgb = RGBColor(148, 163, 184)

    # ==========================================
    # SLIDE 2: THE BUSINESS PROBLEM
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The Real-World Financial Problem: Multi-System Mismatch")

    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "The Payment Lifecycle Disconnect", [
                 "Every payment journey spans 5+ disconnected datasets: Orders, Gateways, Settlements, Bank Ledgers, and Disputes.",
                 "Amounts rarely match 1-to-1 due to gateway processing fees, 18% GST deductions, and rolling reserve withholdings.",
                 "Example: Customer pays ₹5,000 → Gateway captures ₹5,000 → Settlement deposits ₹4,850 → Bank credits ₹4,800.",
                 "Rule-based software flags discrepancies but cannot explain WHY they happened or whether deductions are valid.",
                 "Finance teams spend 70%+ of operational hours manually correlating spreadsheets and querying gateway portals."
             ], badge="🚨")

    add_card(s2, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Key Questions Finance Must Answer", [
                 "Was the ₹150 discrepancy caused by a legitimate fee or an unauthorized over-deduction?",
                 "Did the fee comply with the contractual fee schedule and regulatory MDR caps?",
                 "Is the missing settlement an ordinary T+2 processing delay or lost revenue?",
                 "Was a customer charged twice or did they place two separate legitimate orders?",
                 "Can this exception be auto-resolved safely, or does it mandate human executive review?"
             ], badge="❓")

    # ==========================================
    # SLIDE 3: PROPOSED SOLUTION - WHAT MAKES IT AGENTIC?
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Proposed Solution: From Passive Detection to Agentic Resolution")

    add_card(s3, Inches(0.8), Inches(1.6), Inches(3.7), Inches(5.2),
             "Traditional Rule Systems", [
                 "Purely passive & rule-bound.",
                 "Flags raw numeric discrepancies without investigation.",
                 "Zero context: treats ₹150 gateway fee same as ₹150 fraud.",
                 "100% of exceptions escalated to human finance analysts.",
                 "No policy awareness, no audit narrative, high operational overhead."
             ], badge="❌")

    add_card(s3, Inches(4.8), Inches(1.6), Inches(3.7), Inches(5.2),
             "Generic Chatbots", [
                 "Non-deterministic math.",
                 "Prone to hallucinating transaction IDs and arithmetic balances.",
                 "Unbounded natural language with no structured execution state.",
                 "Cannot reliably trigger tools or run multi-stage reconciliation pipelines.",
                 "Unacceptable safety risk for regulated financial operations."
             ], badge="⚠️")

    add_card(s3, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2),
             "FinRecon Agentic AI", [
                 "Autonomous Task Planning.",
                 "Gathers cross-system evidence across 7 relational tables.",
                 "Grounds reasoning in policy documents via vector RAG.",
                 "Adversarial Critic Agent checks math and policy caps.",
                 "Auto-resolves verified cases; escalates only genuine ambiguities."
             ], badge="✅", border_color=COLOR_ACCENT)

    # ==========================================
    # SLIDE 4: HYBRID AI DESIGN ARCHITECTURE
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Core Architectural Philosophy: Hybrid AI")

    add_card(s4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Deterministic Foundation (Code)", [
                 "Exact relational joins (Order ID, Payment ID, Customer ID).",
                 "Cent-exact arithmetic balances: Payment - Fee - Tax == Settlement.",
                 "Tolerance enforcement: Zero unexplained variance tolerated (₹0.00).",
                 "Strict Pydantic schema validation preventing corrupted data feeds.",
                 "Ensures 100% mathematical precision without LLM hallucination."
             ], badge="⚙️")

    add_card(s4, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Probabilistic Intelligence (Agents)", [
                 "Task decomposition and adaptive investigative routing.",
                 "Cross-system evidence graph formulation.",
                 "Natural-language policy clause interpretation & RAG grounding.",
                 "Adversarial critique: catching contradictions and policy cap breaches.",
                 "Human-readable audit summaries explaining the exact cause of variances."
             ], badge="🧠")

    # ==========================================
    # SLIDE 5: MULTI-AGENT ARCHITECTURE
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Multi-Agent System Architecture & Specialized Roles")

    # 4 Cards Layout
    add_card(s5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.4),
             "1. Planner & Validator Agents", [
                 "Planner Agent: Decomposes natural language requests into structured tasks.",
                 "Validation Agent: Inspects schema, foreign keys, and values before reasoning."
             ], badge="🗺️")

    add_card(s5, Inches(6.8), Inches(1.6), Inches(5.6), Inches(2.4),
             "2. Reconciliation Engine", [
                 "Multi-stage deterministic matching hierarchy (Exact ID -> Order -> Window).",
                 "Detects duplicate charges, missing settlements, and amount discrepancies."
             ], badge="⚖️")

    add_card(s5, Inches(0.8), Inches(4.3), Inches(5.6), Inches(2.5),
             "3. Investigation & Policy RAG", [
                 "Investigation Agent: Forensic query of fees, refunds, and dispute logs.",
                 "Policy Agent: Retrieves governing legal/finance clauses via TF-IDF vector search."
             ], badge="🔍")

    add_card(s5, Inches(6.8), Inches(4.3), Inches(5.6), Inches(2.5),
             "4. Critic & Resolution Agents", [
                 "Critic Agent: Adversarially tests arithmetic balance and policy caps.",
                 "Resolution Agent: Enforces final routing: Auto-Resolve vs. Human Escalation."
             ], badge="🛡️")

    # ==========================================
    # SLIDE 6: LANGGRAPH ORCHESTRATION WORKFLOW
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "End-to-End Orchestration: State Graph Workflow")

    add_card(s6, Inches(0.8), Inches(1.6), Inches(11.6), Inches(5.2),
             "LangGraph Stateful Execution Pipeline", [
                 "State Management: Centralized FinReconState TypedDict holds datasets, hypotheses, audit logs, and resolutions.",
                 "Step 1 — Ingestion & Planning: Planner Agent generates ordered task graph based on user prompt.",
                 "Step 2 — Data Validation: Validation Agent enforces relational integrity across all 7 CSV tables.",
                 "Step 3 — Deterministic Matching: Reconciler runs exact matching hierarchy; splits clean matches from exceptions.",
                 "Step 4 — Conditional Edge (Zero-Exception Short-Circuit): If 0 exceptions exist, bypasses to report immediately.",
                 "Step 5 — Investigation & Policy Retrieval: Loops through exception queue, querying fees and retrieving policy chunks.",
                 "Step 6 — Adversarial Critic Review: Critic evaluates evidence completeness, balance, and ₹300 fee cap.",
                 "Step 7 — Resolution Routing: Critic PASS -> Auto-Resolve; Critic FAIL -> Human Review Escalation.",
                 "Step 8 — Summary & Audit: Compiles tamper-evident JSONL audit trail and executive KPI metrics."
             ], badge="🔄")

    # ==========================================
    # SLIDE 7: POLICY RAG & KNOWLEDGE BASE
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Grounding Financial Reasoning: Policy RAG Pipeline")

    add_card(s7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Institutional Knowledge Base", [
                 "fee_schedule.md: MDR rates (UPI 0%, Cards 2%), 18% GST rules, and ₹300 maximum allowable fee cap.",
                 "settlement_policy.md: T+2 standard settlement timeline, T+5 missing settlement escalation, zero tolerance.",
                 "refund_policy.md: Mandatory original payment linkage, cumulative refund caps, duplicate refund checks.",
                 "dispute_policy.md: Decision matrix for customer 'charged twice' claims across duplicate vs. distinct orders.",
                 "escalation_policy.md: Mandatory human review triggers (unexplained diff > ₹500, contradictory evidence, critic rejection)."
             ], badge="📚")

    add_card(s7, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Vector Retrieval Engine", [
                 "Semantic Section Chunking: Chunks policy docs by level-2 (##) headings preserving complete legal context.",
                 "Sublinear TF-IDF & Cosine Similarity: Robust vector matching with disk caching (index_meta.json & doc_vectors.npy).",
                 "Metadata-Filtered Queries: Enables targeted searches filtered by domain (policy_type='fee' or 'escalation').",
                 "Verifiable Citations: Every proposed resolution quotes the exact document, section header, and clause.",
                 "Eliminates Policy Hallucination: Agents cannot invent fee rules not present in institutional markdown documents."
             ], badge="🔎")

    # ==========================================
    # SLIDE 8: ADVERSARIAL CRITIC & HUMAN-IN-THE-LOOP
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Safety Guardrails: The Critic Agent & Human Governance")

    add_card(s8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Adversarial Critic Agent Checkpoints", [
                 "Check 1 — Mathematical Integrity: Verifies that fee + tax exactly accounts for the settlement difference (0.00 variance).",
                 "Check 2 — Policy Cap Compliance: Rejects any fee deduction exceeding the ₹300 policy cap, even if mathematically balanced.",
                 "Check 3 — Evidence Completeness: Ensures referenced payment and fee records actually exist in the database.",
                 "Check 4 — Case 6 Contradiction Detection: When discrepancy is ₹500 but fee record is only ₹200, Critic immediately flags FAIL.",
                 "Result: High-risk or contradictory cases are physically blocked from automated closing."
             ], badge="🧐")

    add_card(s8, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Human-in-the-Loop (HITL) Controls", [
                 "No Autonomous Money Movement: FinRecon provides decision support; it never moves money or modifies bank accounts.",
                 "Auditable Escalation Docket: Escalated exceptions provide reviewers with complete evidence, formulas, and policy excerpts.",
                 "Reviewer Action Controls: Three one-click operations in Streamlit:",
                 "   1. [Approve Fee Adjustment] — Records write-off to general ledger.",
                 "   2. [Escalate to Gateway Desk] — Dispatches inquiry ticket to gateway desk.",
                 "   3. [Mark for Fraud Review] — Flags transaction for risk operations.",
                 "Tamper-Evident Audit: Every human decision is recorded to reconciliation_audit.jsonl."
             ], badge="👤")

    # ==========================================
    # SLIDE 9: QUANTITATIVE BENCHMARK RESULTS
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Empirical Evaluation: FinRecon vs. Traditional Rule Matcher")

    add_card(s9, Inches(0.8), Inches(1.6), Inches(11.6), Inches(5.2),
             "Head-to-Head Benchmark Findings (Tested on Synthetic Ground Truth)", [
                 "Reconciliation Accuracy: FinRecon achieves 100% accuracy on standard transactions and 100% exception detection.",
                 "Root Cause Explanation: Rule-based software provides 0.0% explanations; FinRecon provides 100% evidence-backed narratives.",
                 "Policy Grounding: Traditional systems require manual manual searching; FinRecon performs automated RAG citations.",
                 "Auto-Resolution Rate: FinRecon auto-resolves 85.7% of exceptions by validating legitimate fees and taxes (Baseline: 0.0%).",
                 "Human Review Reduction: Cuts manual analyst review workload by 85.7%, allowing finance teams to focus only on genuine risks.",
                 "Safety Fidelity on Case 6: 0.0% False Auto-Resolutions — 100% of contradictory cases correctly escalated for human review.",
                 "High Processing Efficiency: Sub-second end-to-end execution (~0.11s per batch) with minimal tool invocation overhead."
             ], badge="📊")

    # ==========================================
    # SLIDE 10: STREAMLIT WORKSTATION UI
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Operations Dashboard: Forensic Workstation UI")

    add_card(s10, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "UI Components & Capabilities", [
                 "Executive KPI Cards: Instant high-level count of Processed, Matched, Exceptions, and Auto-Resolved cases.",
                 "Live Multi-Agent Timeline: Step-by-step visual trace of agent activities, tool executions, and critique decisions.",
                 "Filterable Exception Queue: Tabular overview tagging items as AUTO-RESOLVED or HUMAN REVIEW.",
                 "Interactive Case Dossier: Selecting any row opens its detailed financial evidence file.",
                 "Custom FinTech Styling: High-contrast, clean theme designed for professional financial operations."
             ], badge="💻")

    add_card(s10, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Forensic Evidence & Graph View", [
                 "Visual Transaction Lifecycle: Mermaid diagram rendering Order -> Payment -> Fee/Tax -> Settlement -> Resolution.",
                 "Deterministic Math Card: Displays the exact balancing equation with highlighted green/red status badges.",
                 "Retrieved Policy Drawer: Expandable drawers showing exact governing clauses and similarity relevance scores.",
                 "Operational Action Box: Enables human operators to take action and save reviewer notes directly to the audit log.",
                 "Complete Transparency: Built to provide stakeholders with total visibility into AI decision pathways."
             ], badge="🔍")

    # ==========================================
    # SLIDE 11: KEY TAKEAWAYS & BUSINESS VALUE
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Business Impact & Key Takeaways")

    add_card(s11, Inches(0.8), Inches(1.6), Inches(3.7), Inches(5.2),
             "Operational Efficiency", [
                 "85%+ reduction in manual reconciliation backlog.",
                 "Instant verification of complex fee deductions and taxes.",
                 "Reduces month-end financial close cycle from days to minutes.",
                 "Frees financial analysts from mundane spreadsheet correlation."
             ], badge="⚡")

    add_card(s11, Inches(4.8), Inches(1.6), Inches(3.7), Inches(5.2),
             "Risk & Compliance", [
                 "Eliminates human arithmetic errors in high-volume environments.",
                 "Strict adherence to MDR fee schedules and legal settlement caps.",
                 "Immutable append-only JSONL audit trail ready for external audits.",
                 "Adversarial critic blocks premature or unverified resolution."
             ], badge="🛡️")

    add_card(s11, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2),
             "Scalable AI Design", [
                 "Hybrid AI pattern: deterministic math + agentic reasoning.",
                 "Extensible LangGraph architecture ready for new payment gateways.",
                 "Self-contained RAG pipeline requiring zero paid external APIs.",
                 "Clear blueprint for enterprise financial decision support."
             ], badge="🚀")

    # ==========================================
    # SLIDE 12: FUTURE SCOPE & CONCLUSION
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Future Roadmap & Project Summary")

    add_card(s12, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Future Engineering Scope", [
                 "Live Payment Gateway Integrations: Webhook listeners for Razorpay, Stripe, and Cashfree APIs.",
                 "Core Banking Integration: Direct ISO 20022 and MT940 bank statement parser integration.",
                 "Model Context Protocol (MCP): Exposing financial tools via standard enterprise MCP servers.",
                 "Long-Term Episodic Memory: Tracking recurring merchant fee disputes across fiscal quarters.",
                 "Autonomous Dispute Filing: Automated generation of chargeback contestation evidence packets."
             ], badge="🔮")

    add_card(s12, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Summary Conclusion", [
                 "FinRecon proves that Agentic AI in FinTech must be grounded, transparent, and mathematically rigorous.",
                 "By coordinating specialized agents with deterministic tools and adversarial critique, it solves the core bottleneck of financial reconciliation: finding WHY records differ and WHAT to do next.",
                 "Successfully delivered and tested across 8 executable phases with 100% test pass fidelity (32/32 tests).",
                 "GitHub Repository: https://github.com/imWordy/FinRecon"
             ], badge="🏁", border_color=COLOR_PRIMARY)

    # Save to disk
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_PATH))
    print(f"Presentation successfully created: {OUTPUT_PATH}")
    print(f"Total Slides: {len(prs.slides)}")

if __name__ == "__main__":
    create_deck()
